"""
SmartFeeder EMS — Feeder Detailed Demand Analysis Automated Test Suite
Verifies Feeder-1 Detailed Analysis page, reusable architecture for Feeders 1-8,
date availability rules, peak demand duration, dynamic Excel download,
and forecast abstraction.
"""
import unittest
import json
import os
import sys
import io
import openpyxl

# Ensure root directory is on python path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.app import create_app
from backend.config import Config
from backend.feeder_service import FeederService
from backend.forecast_service import ForecastService

class TestFeederAnalysis(unittest.TestCase):
    def setUp(self):
        self.app = create_app(Config)
        self.app.config["TESTING"] = True
        self.client = self.app.test_client()
        self.admin_username = Config.ADMIN_USERNAME
        self.admin_password = Config.ADMIN_PASSWORD

    def _login(self):
        """Helper to establish an authenticated session."""
        return self.client.post("/api/auth/login",
                                data=json.dumps({"username": self.admin_username, "password": self.admin_password}),
                                content_type="application/json")

    def test_01_unauthenticated_feeder_redirects_to_login(self):
        """TEST 1: Unauthenticated access to /feeder/1 redirects to /login."""
        res = self.client.get("/feeder/1", follow_redirects=False)
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login", res.headers.get("Location", ""))

    def test_02_authenticated_feeder1_renders_successfully(self):
        """TEST 2: Authenticated access renders Feeder-1 analysis page with required elements."""
        self._login()
        res = self.client.get("/feeder/1")
        self.assertEqual(res.status_code, 200)
        html = res.data.decode("utf-8")

        # Headings & Subtitles
        self.assertIn("FEEDER-1", html)
        self.assertIn("Residential Feeder", html)
        self.assertIn("11 kV Feeder Demand Analysis", html)
        self.assertIn("LIVE", html)

        # Sequential Navigation & Breadcrumb
        self.assertIn("← Back to BUS-1", html)
        self.assertIn('href="/bus/1"', html)
        self.assertIn("BUS-1", html)

        # Control Bar
        self.assertIn("feederDatePicker", html)
        self.assertIn("Actual Demand vs Time", html)
        self.assertIn("Forecasted Demand vs Time", html)
        self.assertIn("Actual vs Forecasted Demand", html)
        self.assertIn("Download Excel", html)

        # Primary Chart & Empty State Containers
        self.assertIn("feederDemandChart", html)
        self.assertIn("DATA NOT AVAILABLE", html)

        # Daily Summary Section
        self.assertIn("PEAK DEMAND", html)
        self.assertIn("PEAK DEMAND TIME", html)
        self.assertIn("PEAK DEMAND DURATION", html)
        self.assertIn("DAILY ENERGY", html)

        # Footer Credit
        self.assertIn("Team Nexora", html)

    def test_03_telemetry_for_latest_date_2026_09_29(self):
        """TEST 3: API returns full 24-hour telemetry on valid dataset date (2026-09-29)."""
        self._login()
        res = self.client.get("/api/feeder/1/data?date=2026-09-29")
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data["success"])
        data = json_data["data"]

        self.assertTrue(data["available"])
        self.assertEqual(data["feeder"]["name"], "FEEDER-1")
        self.assertEqual(len(data["hourly"]), 24)

        # Summary Metrics
        summary = data["summary"]
        self.assertAlmostEqual(summary["peak_demand_kw"], 45.44, delta=0.1)
        self.assertEqual(summary["peak_time"], "15:00")
        self.assertGreaterEqual(summary["peak_duration_hours"], 1)
        self.assertAlmostEqual(summary["daily_energy_kwh"], 689.8, delta=1.0)

    def test_04_data_unavailable_before_05_aug_2026(self):
        """TEST 4: Selecting date before 05-Aug-2026 returns available=False with exact guidance."""
        self._login()
        res = self.client.get("/api/feeder/1/data?date=2026-08-04")
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data["success"])
        data = json_data["data"]

        self.assertFalse(data["available"])
        self.assertEqual(data["message"], "DATA NOT AVAILABLE")
        self.assertIn("Actual feeder data is available from 05-Aug-2026 onwards.", data["info"])

    def test_05_data_available_on_start_date_05_aug_2026(self):
        """TEST 5: Selecting 05-Aug-2026 returns available=True with valid records."""
        self._login()
        res = self.client.get("/api/feeder/1/data?date=2026-08-05")
        self.assertEqual(res.status_code, 200)
        json_data = res.get_json()
        self.assertTrue(json_data["success"])
        data = json_data["data"]

        self.assertTrue(data["available"])
        self.assertEqual(len(data["hourly"]), 24)
        self.assertGreater(data["summary"]["peak_demand_kw"], 0)

    def test_06_excel_download_generation(self):
        """TEST 6: /api/feeder/1/download generates valid Excel workbook with hourly and summary sheets."""
        self._login()
        res = self.client.get("/api/feeder/1/download?date=2026-09-29")
        self.assertEqual(res.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res.headers.get("Content-Type", ""))
        self.assertIn("SmartFeeder_Feeder1_2026-09-29.xlsx", res.headers.get("Content-Disposition", ""))

        # Verify workbook content
        wb = openpyxl.load_workbook(io.BytesIO(res.data))
        self.assertIn("Hourly Telemetry", wb.sheetnames)
        self.assertIn("Daily Summary", wb.sheetnames)

        ws_hourly = wb["Hourly Telemetry"]
        self.assertEqual(ws_hourly["A4"].value, "Date")
        self.assertEqual(ws_hourly["B4"].value, "Time")
        self.assertEqual(ws_hourly["E4"].value, "Actual Demand (kW)")
        self.assertEqual(ws_hourly["F4"].value, "Forecast Demand (kW)")

        # Verify 24 data rows exist
        self.assertEqual(ws_hourly.max_row, 28)

        ws_summary = wb["Daily Summary"]
        self.assertEqual(ws_summary["A4"].value, "Feeder Designation")
        self.assertEqual(ws_summary["B4"].value, "FEEDER-1")

    def test_07_forecast_service_determinism(self):
        """TEST 7: ForecastService returns deterministic, technically plausible load profile."""
        f1 = ForecastService.get_hourly_forecast(1, "2026-09-29", 14, 40.0)
        f2 = ForecastService.get_hourly_forecast(1, "2026-09-29", 14, 40.0)
        self.assertEqual(f1, f2)
        # Stays within realistic error bound
        self.assertGreaterEqual(f1, 37.0)
        self.assertLessEqual(f1, 43.0)

    def test_08_reusable_architecture_feeders_1_to_8(self):
        """TEST 8: Reusable template supports Feeders 1 to 8 with correct Bus interconnections."""
        self._login()
        # Feeder 1 on BUS-1
        res1 = self.client.get("/feeder/1")
        self.assertEqual(res1.status_code, 200)
        html1 = res1.data.decode("utf-8")
        self.assertIn("← Back to BUS-1", html1)
        self.assertIn('href="/bus/1"', html1)
        self.assertIn("FEEDER-1", html1)

        # Feeder 5 on BUS-2
        res5 = self.client.get("/feeder/5")
        self.assertEqual(res5.status_code, 200)
        html5 = res5.data.decode("utf-8")
        self.assertIn("← Back to BUS-2", html5)
        self.assertIn('href="/bus/2"', html5)
        self.assertIn("FEEDER-5", html5)

    def test_09_bus1_and_bus2_feeder_classifications_consistency(self):
        """TEST 9: Verify BUS-1 and BUS-2 display exact approved feeder classifications."""
        self._login()

        # BUS-1
        res_b1 = self.client.get("/bus/1")
        self.assertEqual(res_b1.status_code, 200)
        html_b1 = res_b1.data.decode("utf-8")
        self.assertIn("FEEDER-1", html_b1)
        self.assertIn("Residential Feeder", html_b1)
        self.assertIn("FEEDER-2", html_b1)
        self.assertIn("Commercial Feeder", html_b1)
        self.assertIn("FEEDER-3", html_b1)
        self.assertIn("Industrial Feeder", html_b1)
        self.assertIn("FEEDER-4", html_b1)
        self.assertIn("Mixed / Critical Load", html_b1)

        # BUS-2
        res_b2 = self.client.get("/bus/2")
        self.assertEqual(res_b2.status_code, 200)
        html_b2 = res_b2.data.decode("utf-8")
        self.assertIn("FEEDER-5", html_b2)
        self.assertIn("Commercial Feeder", html_b2)
        self.assertIn("FEEDER-6", html_b2)
        self.assertIn("Rural / Agricultural Feeder", html_b2)
        self.assertIn("FEEDER-7", html_b2)
        self.assertIn("Industrial / Manufacturing Feeder", html_b2)
        self.assertIn("FEEDER-8", html_b2)
        self.assertIn("Residential / Municipal Feeder", html_b2)

    def test_10_all_eight_feeder_pages_exact_metadata(self):
        """TEST 10: Verify all 8 individual feeder detail pages display exact name, classification, and Data ID."""
        self._login()
        expected = [
            (1, "FEEDER-1", "Residential Feeder", "ID 107", "BUS-1"),
            (2, "FEEDER-2", "Commercial Feeder", "ID 119", "BUS-1"),
            (3, "FEEDER-3", "Industrial Feeder", "ID 129", "BUS-1"),
            (4, "FEEDER-4", "Mixed / Critical Load", "ID 135", "BUS-1"),
            (5, "FEEDER-5", "Commercial Feeder", "ID 145", "BUS-2"),
            (6, "FEEDER-6", "Rural / Agricultural Feeder", "ID 152", "BUS-2"),
            (7, "FEEDER-7", "Industrial / Manufacturing Feeder", "ID 155", "BUS-2"),
            (8, "FEEDER-8", "Residential / Municipal Feeder", "ID 157", "BUS-2"),
        ]

        for fid, name, classification, data_id, bus in expected:
            res = self.client.get(f"/feeder/{fid}")
            self.assertEqual(res.status_code, 200, f"Feeder {fid} should load successfully")
            html = res.data.decode("utf-8")
            self.assertIn(name, html, f"{name} should be present in page")
            self.assertIn(classification, html, f"{classification} should be present in page")
            self.assertIn(data_id, html, f"{data_id} should be present in page")
            import re
            plain_text = " ".join(re.sub(r"<[^>]+>", " ", html).split())
            self.assertIn(f"Home / {bus} / {name}", plain_text, f"Breadcrumb should be consistent")

if __name__ == "__main__":
    unittest.main()
